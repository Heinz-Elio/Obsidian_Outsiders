from __future__ import annotations

import argparse
import json
import re
import unicodedata
from collections import OrderedDict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from app.config import PROJECT_ROOT

OUTPUT_SUBDIR = Path("IU7") / "00_大綱索引"
DEFAULT_SOURCE = PROJECT_ROOT / "Setting" / "大綱及行動.xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "Setting" / OUTPUT_SUBDIR
BACKGROUND_TAG = " #背景"
FIRST_APPEARANCE_TAG = " #初登場"

BACKGROUND_RGB = {"CCCCCC", "D9D9D9"}
DEBUT_RGB = {"FFFF00"}
INVALID_FILENAME = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
TIME_VALUE = re.compile(r"^\s*\d{1,2}\s*[,，/]\s*\d{1,2}:\d{2}\s*$")


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def _frontmatter_text(value: object) -> str:
    return json.dumps(value, ensure_ascii=False)


def _english_filename(value: str, *, fallback: str = "untitled") -> str:
    ascii_value = (
        unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    )
    cleaned = INVALID_FILENAME.sub("_", ascii_value)
    cleaned = re.sub(r"[^A-Za-z0-9]+", "_", cleaned).strip("_").casefold()
    return cleaned[:100] or fallback


def _identifier(value: str) -> str:
    cleaned = re.sub(r"[^\w]+", "_", value.casefold(), flags=re.UNICODE).strip("_")
    return f"outline_{cleaned or 'untitled'}"


def _cell_rgb(cell: Cell) -> str | None:
    if cell.fill.fill_type != "solid":
        return None
    color = cell.fill.fgColor
    if color.type != "rgb" or not isinstance(color.rgb, str):
        return None
    return color.rgb[-6:].upper()


def _is_background(cell: Cell) -> bool:
    return _cell_rgb(cell) in BACKGROUND_RGB


def _is_debut(cell: Cell) -> bool:
    return _cell_rgb(cell) in DEBUT_RGB


def _last_used_row(sheet: Worksheet) -> int:
    return max(
        (cell.row for row in sheet.iter_rows() for cell in row if cell.value is not None),
        default=0,
    )


def _last_used_column(sheet: Worksheet) -> int:
    return max(
        (
            cell.column
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        ),
        default=0,
    )


@dataclass
class SheetView:
    values: Worksheet
    formulas: Worksheet

    def __post_init__(self) -> None:
        self._merged: dict[tuple[int, int], tuple[int, int]] = {}
        for merged in self.formulas.merged_cells.ranges:
            anchor = (merged.min_row, merged.min_col)
            for row in range(merged.min_row, merged.max_row + 1):
                for column in range(merged.min_col, merged.max_col + 1):
                    self._merged[(row, column)] = anchor

    @property
    def title(self) -> str:
        return self.formulas.title

    @property
    def max_row(self) -> int:
        return max(_last_used_row(self.values), _last_used_row(self.formulas))

    @property
    def max_column(self) -> int:
        return max(_last_used_column(self.values), _last_used_column(self.formulas))

    def value(self, row: int, column: int, *, merged: bool = False) -> object:
        if merged and (row, column) in self._merged:
            row, column = self._merged[(row, column)]

        cached = self.values.cell(row, column).value
        if cached is not None:
            return cached
        return self.formulas.cell(row, column).value

    def cell(self, row: int, column: int) -> Cell:
        return self.formulas.cell(row, column)


def _document_header(
    title: str,
    document_id: str,
    source: Path,
    sheet: str,
    *,
    operation: str | None = None,
) -> list[str]:
    lines = [
        "---",
        "type: outline",
        f"id: {_frontmatter_text(document_id)}",
        "timeline: iu7",
        f"source_file: {_frontmatter_text(source.name)}",
        f"source_sheet: {_frontmatter_text(sheet)}",
    ]
    if operation:
        lines.append(f"operation: {_frontmatter_text(operation)}")
    lines.extend(
        [
            "generated: true",
            "---",
            f"# {title}",
            "",
            "<!-- 此檔由 app/outline_converter.py 自動產生，請勿手動編輯。 -->",
            "",
        ]
    )
    return lines


def _integer(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and float(value).is_integer():
        return int(value)
    text = _text(value)
    return int(text) if text.isdigit() else None


def _date_label(year: object, month: object, day: object) -> str:
    numeric = tuple(_integer(value) for value in (year, month, day))
    if all(value is not None for value in numeric):
        return f"{numeric[0]:04d}-{numeric[1]:02d}-{numeric[2]:02d}"

    year_text = _text(year)
    month_text = _text(month)
    day_text = _text(day)
    parts: list[str] = []
    if year_text:
        parts.append(year_text if year_text.endswith("年") else f"{year_text}年")
    if month_text:
        parts.append(month_text if month_text.endswith("月") else f"{month_text}月")
    if day_text:
        parts.append(day_text if day_text.endswith("日") else f"{day_text}日")
    return "".join(parts) or "日期未定"


def _month_label(year: object, month: object) -> str:
    numeric_year = _integer(year)
    numeric_month = _integer(month)
    if numeric_year is not None and numeric_month is not None:
        return f"{numeric_year:04d}-{numeric_month:02d}"

    year_text = _text(year)
    month_text = _text(month)
    if year_text and month_text:
        return f"{year_text}年{month_text}月"
    return month_text or year_text or "月份未定"


def _render_outline(sheet: SheetView, source: Path) -> dict[str, str]:
    groups: OrderedDict[str, list[dict[str, object]]] = OrderedDict()

    for row in range(2, sheet.max_row + 1):
        event = _text(sheet.value(row, 6))
        if not event:
            continue
        arc = _text(sheet.value(row, 1, merged=True)) or "未分篇"
        groups.setdefault(arc, []).append(
            {
                "event": event,
                "question": _text(sheet.value(row, 2, merged=True)),
                "date": _date_label(
                    sheet.value(row, 3, merged=True),
                    sheet.value(row, 4, merged=True),
                    sheet.value(row, 5, merged=True),
                ),
                "episode": _text(sheet.value(row, 7, merged=True)),
                "background": _is_background(sheet.cell(row, 6)),
            }
        )

    documents: dict[str, str] = {}
    for arc, entries in groups.items():
        title = f"故事大綱｜{arc}"
        lines = _document_header(
            title,
            _identifier(arc),
            source,
            sheet.title,
            operation=arc if arc != "未分篇" else None,
        )
        lines.append(f"{BACKGROUND_TAG}=背景事件，不細寫。")
        current_date: object = None
        for entry in entries:
            if entry["date"] != current_date:
                current_date = entry["date"]
                lines.extend(["", f"## {current_date}", ""])

            first, *rest = str(entry["event"]).split("\n")
            tag = BACKGROUND_TAG if entry["background"] else ""
            lines.append(f"- {first}{tag}")
            lines.extend(f"  {line}" for line in rest if line.strip())
            if entry["question"]:
                question = str(entry["question"]).replace("\n", "；")
                lines.append(f"  - 懸念／功能: {question}")
            if entry["episode"]:
                lines.append(f"  - 對應集數: {entry['episode']}")

        english_arc = (
            "unassigned"
            if arc == "未分篇"
            else _english_filename(arc, fallback="unnamed_arc")
        )
        filename = f"outline_{english_arc}.md"
        documents[filename] = "\n".join(lines).rstrip() + "\n"
    return documents


def _render_transformation_kpi(sheet: SheetView, source: Path) -> dict[str, str]:
    monthly_kpi: list[tuple[str, str]] = []
    covered_rows: set[int] = set()

    ranges = sorted(
        (
            merged
            for merged in sheet.formulas.merged_cells.ranges
            if merged.min_col == 9 and merged.max_col == 9
        ),
        key=lambda merged: merged.min_row,
    )
    for merged in ranges:
        covered_rows.update(range(merged.min_row, merged.max_row + 1))
        kpi = _text(sheet.value(merged.min_row, 9))
        months = list(
            OrderedDict.fromkeys(
                _month_label(
                    sheet.value(row, 3, merged=True),
                    sheet.value(row, 4, merged=True),
                )
                for row in range(merged.min_row, merged.max_row + 1)
                if _text(sheet.value(row, 4, merged=True))
            )
        )
        month = "／".join(months) if months else "月份未定"
        if kpi:
            monthly_kpi.append((month, kpi))

    for row in range(2, sheet.max_row + 1):
        if row in covered_rows:
            continue
        kpi = _text(sheet.value(row, 9))
        if not kpi:
            continue
        monthly_kpi.append(
            (
                _month_label(
                    sheet.value(row, 3, merged=True),
                    sheet.value(row, 4, merged=True),
                ),
                kpi,
            )
        )

    transformations: list[tuple[str, str, str, str]] = []
    for row in range(2, sheet.max_row + 1):
        count = _text(sheet.value(row, 8))
        if not count:
            continue
        transformations.append(
            (
                _month_label(
                    sheet.value(row, 3, merged=True),
                    sheet.value(row, 4, merged=True),
                ),
                _date_label(
                    sheet.value(row, 3, merged=True),
                    sheet.value(row, 4, merged=True),
                    sheet.value(row, 5, merged=True),
                ),
                _text(sheet.value(row, 1, merged=True)) or "未分篇",
                count,
            )
        )

    lines = _document_header(
        "變身人數與水妖每月KPI",
        "outline_transformation_kpi",
        source,
        sheet.title,
    )
    lines.extend(
        [
            "## 每月水妖KPI",
            "",
            "| 月份 | 水妖KPI |",
            "| --- | ---: |",
        ]
    )
    for month, kpi in monthly_kpi:
        lines.append(f"| {_markdown_cell(month)} | {_markdown_cell(kpi)} |")

    lines.extend(
        [
            "",
            "## 變身紀錄",
            "",
            "| 月份 | 日期 | 篇章 | 變身人數 |",
            "| --- | --- | --- | ---: |",
        ]
    )
    for record in transformations:
        lines.append(f"| {' | '.join(_markdown_cell(value) for value in record)} |")

    return {"transformation_kpi.md": "\n".join(lines).rstrip() + "\n"}


def _render_water_demons(sheet: SheetView, source: Path) -> dict[str, str]:
    characters = {
        column: _text(sheet.value(1, column))
        for column in range(2, 7)
        if _text(sheet.value(1, column))
    }
    groups: OrderedDict[str, list[dict[str, object]]] = OrderedDict()
    current_arc = "未分篇"

    for row in range(2, sheet.max_row + 1):
        explicit_arc = _text(sheet.value(row, 1))
        if explicit_arc:
            current_arc = explicit_arc

        actions: OrderedDict[str, list[dict[str, object]]] = OrderedDict()
        for column, character in characters.items():
            action = _text(sheet.value(row, column))
            if not action:
                continue
            cell = sheet.cell(row, column)
            actions.setdefault(action, []).append(
                {
                    "character": character,
                    "debut": _is_debut(cell),
                    "background": _is_background(cell),
                }
            )

        for action, participants in actions.items():
            groups.setdefault(current_arc, []).append(
                {
                    "row": row,
                    "episode": _text(sheet.value(row, 7)),
                    "action": action,
                    "participants": participants,
                    "background": all(
                        bool(participant["background"]) for participant in participants
                    ),
                }
            )

    lines = _document_header(
        "水妖行動與初登場",
        "outline_water_demon_actions",
        source,
        sheet.title,
    )
    for arc, actions in groups.items():
        lines.extend(
            ["", f"## {arc}", "", "| 行動 | 原著集數 | 參與水妖 |", "| --- | ---: | --- |"]
        )
        for action in actions:
            participants = "、".join(
                f"{participant['character']}"
                f"{FIRST_APPEARANCE_TAG if participant['debut'] else ''}"
                f"{BACKGROUND_TAG if participant['background'] and not action['background'] else ''}"
                for participant in action["participants"]
            )
            name = f"{action['action']}{BACKGROUND_TAG if action['background'] else ''}"
            lines.append(
                f"| {_markdown_cell(name)} | {_markdown_cell(action['episode'])} "
                f"| {_markdown_cell(participants)} |"
            )

    return {"water_demon_actions.md": "\n".join(lines).rstrip() + "\n"}


def _looks_like_time(value: object) -> bool:
    return bool(TIME_VALUE.match(_text(value)))


def _operation_columns(sheet: SheetView) -> list[tuple[int, int, str]]:
    number_columns = [
        column
        for column in range(1, sheet.max_column + 1)
        if isinstance(sheet.value(1, column), (int, float))
        and not isinstance(sheet.value(1, column), bool)
    ]
    operations: list[tuple[int, int, str]] = []
    for index, value_column in enumerate(number_columns):
        start = value_column
        if value_column > 1 and _text(sheet.value(1, value_column - 1)).casefold().startswith(
            "operation"
        ):
            start = value_column - 1
        end = (
            number_columns[index + 1] - 1
            if index + 1 < len(number_columns)
            else sheet.max_column
        )
        operations.append((start, end, _text(sheet.value(1, value_column))))
    return operations


def _personnel(sheet: SheetView, value_column: int, end_column: int) -> str:
    first = _text(sheet.value(3, value_column))
    second = _text(sheet.value(3, value_column + 1)) if value_column < end_column else ""
    if first and not second:
        return first

    teams: list[str] = []
    for row in (3, 4):
        team = _text(sheet.value(row, value_column))
        members = (
            _text(sheet.value(row, value_column + 1))
            if value_column < end_column
            else ""
        )
        if team and members:
            teams.append(f"{team}組：{members}")
        elif team:
            teams.append(team)
        elif members:
            teams.append(members)
    return "；".join(teams)


def _teams(sheet: SheetView, value_column: int, end_column: int) -> set[str]:
    if value_column >= end_column:
        return set()
    return {
        _text(sheet.value(row, value_column))
        for row in (3, 4)
        if _text(sheet.value(row, value_column))
        and _text(sheet.value(row, value_column + 1))
    }


@dataclass
class Lane:
    heading: str
    time_column: int
    event_column: int


@dataclass
class TimelineBlock:
    lanes: list[Lane]
    first_row: int
    last_row: int


def _timeline_blocks(sheet: SheetView) -> list[TimelineBlock]:
    """Rows with merged two-column headings start a block of side-by-side lanes.

    Rows inside a block are aligned by absolute time; each lane keeps local time.
    """
    headings: dict[int, list[tuple[int, str]]] = {}
    for merged in sheet.formulas.merged_cells.ranges:
        if merged.min_row != merged.max_row or merged.max_col != merged.min_col + 1:
            continue
        heading = _text(sheet.value(merged.min_row, merged.min_col))
        if heading:
            headings.setdefault(merged.min_row, []).append((merged.min_col, heading))

    rows = sorted(headings)
    blocks: list[TimelineBlock] = []
    for index, heading_row in enumerate(rows):
        first_row = heading_row + 1
        last_row = rows[index + 1] - 1 if index + 1 < len(rows) else sheet.max_row
        lanes: list[Lane] = []
        for column, heading in sorted(headings[heading_row]):
            left_score, right_score = (
                sum(
                    _looks_like_time(sheet.value(row, candidate))
                    for row in range(first_row, last_row + 1)
                )
                for candidate in (column, column + 1)
            )
            if left_score > right_score:
                lanes.append(Lane(heading, column, column + 1))
            else:
                lanes.append(Lane(heading, column + 1, column))
        blocks.append(TimelineBlock(lanes, first_row, last_row))
    return blocks


def _time_label(value: object) -> str:
    text = _text(value)
    match = re.match(r"^\s*(\d{1,2})\s*[,，/]\s*(\d{1,2}:\d{2})\s*$", text)
    return f"{match[1]}日 {match[2]}" if match else text


def _render_timeline(sheet: SheetView, block: TimelineBlock) -> list[str]:
    lane_columns = {
        column for lane in block.lanes for column in (lane.time_column, lane.event_column)
    }
    rows: list[tuple[list[str], str]] = []
    for row in range(block.first_row, block.last_row + 1):
        cells: list[str] = []
        for lane in block.lanes:
            cells.append(_time_label(sheet.value(row, lane.time_column)))
            cells.append(_text(sheet.value(row, lane.event_column)))
        note = "；".join(
            _text(sheet.value(row, column))
            for column in range(1, sheet.max_column + 1)
            if column not in lane_columns and _text(sheet.value(row, column))
        )
        if any(cells) or note:
            rows.append((cells, note))

    has_notes = any(note for _, note in rows)
    header = [part for lane in block.lanes for part in (f"{lane.heading}時間", lane.heading)]
    if has_notes:
        header.append("備註")
    lines = [
        f"## {' × '.join(lane.heading for lane in block.lanes)}",
        "",
        f"| {' | '.join(header)} |",
        f"| {' | '.join('---' for _ in header)} |",
    ]
    for cells, note in rows:
        values = [*cells, note] if has_notes else cells
        lines.append(f"| {' | '.join(_markdown_cell(value) for value in values)} |")
    lines.append("")
    return lines


def _render_actions(sheet: SheetView, source: Path) -> dict[str, str]:
    documents: dict[str, str] = {}
    labels = {
        2: "日期",
        3: "行動人員",
        5: "通訊網絡",
        6: "目標",
        7: "目的",
    }

    operations = [
        (
            start_column,
            end_column,
            operation,
            start_column + 1 if start_column < end_column and _text(
                sheet.value(1, start_column)
            ).casefold().startswith("operation") else start_column,
        )
        for start_column, end_column, operation in _operation_columns(sheet)
    ]
    assigned: dict[str, list[TimelineBlock]] = {operation: [] for _, _, operation, _ in operations}
    for block in _timeline_blocks(sheet):
        headings = {lane.heading for lane in block.lanes}
        first_column = min(lane.time_column for lane in block.lanes)
        owner = next(
            (
                operation
                for _, end_column, operation, value_column in operations
                if headings <= _teams(sheet, value_column, end_column)
            ),
            next(
                (
                    operation
                    for start_column, end_column, operation, _ in operations
                    if start_column <= first_column <= end_column
                ),
                None,
            ),
        )
        if owner is not None:
            assigned[owner].append(block)

    for start_column, end_column, operation, value_column in operations:
        title = f"行動規劃｜Operation {operation}"
        lines = _document_header(
            title,
            _identifier(f"operation_{operation}"),
            source,
            sheet.title,
            operation=f"Operation {operation}",
        )
        lines.extend(["## 基本資料", ""])
        for row, label in labels.items():
            value = (
                _personnel(sheet, value_column, end_column)
                if row == 3
                else _text(sheet.value(row, value_column))
            )
            if value:
                lines.append(f"- {label}: {value}")
        lines.append("")

        if assigned[operation]:
            lines.extend(["同一列＝同一時刻；各欄時間為該組當地時間。", ""])
        for block in assigned[operation]:
            lines.extend(_render_timeline(sheet, block))

        filename = f"operation_{int(float(operation)):02d}.md"
        documents[filename] = "\n".join(lines).rstrip() + "\n"
    return documents


def _markdown_cell(value: object) -> str:
    return _text(value).replace("|", "\\|").replace("\n", "<br>")


def _render_expenses(sheet: SheetView, source: Path) -> dict[str, str]:
    lines = _document_header(
        "行動開支",
        "outline_action_expenses",
        source,
        sheet.title,
    )
    lines.extend(
        [
            "幣別：原工作表未註明",
            "",
            "| 類別 | 項目／路線 | 等級 | 單價 | 數量 | 總額 |",
            "| --- | --- | --- | ---: | ---: | ---: |",
        ]
    )

    for row in range(2, sheet.max_row + 1):
        values = [sheet.value(row, column, merged=True) for column in range(1, 7)]
        if not any(value is not None for value in values):
            continue
        if row == 10:
            values[1] = "機票小計"
        elif row == sheet.max_row and not any(values[:5]):
            values[1] = "總計"
        lines.append(f"| {' | '.join(_markdown_cell(value) for value in values)} |")

    return {"action_expenses.md": "\n".join(lines).rstrip() + "\n"}


def convert_workbook(source: Path) -> dict[str, str]:
    source = source.resolve()
    value_book = load_workbook(source, data_only=True)
    formula_book = load_workbook(source, data_only=False)
    required = ("大綱", "水妖", "行動", "行動開支")
    missing = [name for name in required if name not in formula_book.sheetnames]
    if missing:
        raise ValueError(f"workbook is missing required sheets: {', '.join(missing)}")

    documents: dict[str, str] = {}
    outline = SheetView(value_book["大綱"], formula_book["大綱"])
    documents.update(_render_outline(outline, source))
    documents.update(_render_transformation_kpi(outline, source))

    renderers = {
        "水妖": _render_water_demons,
        "行動": _render_actions,
        "行動開支": _render_expenses,
    }
    for name, renderer in renderers.items():
        sheet = SheetView(value_book[name], formula_book[name])
        documents.update(renderer(sheet, source))
    return {
        f"{index:02d}_{filename}": content
        for index, (filename, content) in enumerate(documents.items(), start=1)
    }


def write_documents(
    documents: dict[str, str],
    output: Path,
    *,
    check: bool = False,
) -> bool:
    existing = {
        path.name: path.read_text(encoding="utf-8")
        for path in output.glob("*.md")
    } if output.exists() else {}
    matches = existing == documents
    if check:
        return matches

    output.mkdir(parents=True, exist_ok=True)
    for stale in set(existing) - set(documents):
        (output / stale).unlink()
    for filename, content in documents.items():
        path = output / filename
        if existing.get(filename) == content:
            continue
        with path.open("w", encoding="utf-8", newline="\n") as handle:
            handle.write(content)
    return matches


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Convert the outline workbook into Markdown sources for RAG."
    )
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--check",
        action="store_true",
        help="Exit with an error when generated Markdown is stale.",
    )
    args = parser.parse_args()

    documents = convert_workbook(args.source)
    matches = write_documents(documents, args.output, check=args.check)
    if args.check and not matches:
        raise SystemExit("generated outline Markdown is stale")
    if args.check:
        print(f"outline Markdown is current ({len(documents)} files)")
    else:
        status = "unchanged" if matches else "updated"
        print(f"{status} {len(documents)} outline Markdown files in {args.output}")


if __name__ == "__main__":
    main()
