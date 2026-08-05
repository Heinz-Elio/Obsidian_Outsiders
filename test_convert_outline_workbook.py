from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill

import convert_outline_workbook
from app.indexer import _refresh_outline_sources


GRAY = PatternFill(fill_type="solid", fgColor="D9D9D9")
YELLOW = PatternFill(fill_type="solid", fgColor="FFFF00")


def _generated(documents: dict[str, str], suffix: str) -> str:
    return next(content for name, content in documents.items() if name.endswith(suffix))


def _workbook(path: Path) -> Path:
    workbook = Workbook()

    outline = workbook.active
    outline.title = "大綱"
    outline.append(["篇章", "懸念", "年", "月", "日", "事件", "ep", "ct", "Ttl"])
    outline.merge_cells("A2:A3")
    outline.merge_cells("C2:C3")
    outline.merge_cells("I2:I3")
    outline["A2"] = "Operation Test"
    outline["C2"] = 2004
    outline["B2"] = "核心問題"
    outline["D2"] = 4
    outline["E2"] = 5
    outline["F2"] = "主要事件"
    outline["G2"] = 1
    outline["H2"] = 2
    outline["I2"] = 3
    outline["D3"] = 4
    outline["E3"] = 6
    outline["F3"] = "背景事件"
    outline["F3"].fill = GRAY
    outline["F4"] = "未分篇事件"

    demons = workbook.create_sheet("水妖")
    demons["B1"] = "BBS"
    demons["C1"] = "LB"
    demons["A2"] = "Siskin"
    demons["B2"] = "星羅誕生"
    demons["B2"].fill = YELLOW
    demons["G2"] = 7
    demons["C3"] = "冰塊行動"
    demons["A4"] = "Stargazer"
    demons["B4"] = "背景行動"
    demons["B4"].fill = GRAY

    actions = workbook.create_sheet("行動")
    actions["A1"] = "Operation "
    actions["B1"] = 1
    actions["A2"] = "日期"
    actions["B2"] = "4月5-7日"
    actions["A3"] = "行動人員"
    actions["B3"] = "Kafziel, Soar"
    actions["A5"] = "通訊網絡"
    actions["B5"] = "C4"
    actions["A6"] = "目標"
    actions["B6"] = "米迦勒"
    actions["A7"] = "目的"
    actions["B7"] = "確認身份"
    actions.merge_cells("A9:B9")
    actions["A9"] = "印度"
    actions["A10"] = "離開檀香山"
    actions["B10"] = "2, 11:00"

    expenses = workbook.create_sheet("行動開支")
    expenses.append([None, None, None, "UNIT", "QUANTITY", "TOTAL"])
    expenses.append(["機票", "夏威夷-科契", "PREM", 1790, 2, 3580])
    expenses.append([None, None, None, None, None, 3580])

    workbook.save(path)
    return path


def test_outline_uses_merged_context_without_filling_normal_blanks(tmp_path):
    documents = convert_outline_workbook.convert_workbook(
        _workbook(tmp_path / "outline.xlsx")
    )

    assigned = _generated(documents, "outline_operation_test.md")
    unassigned = _generated(documents, "outline_unassigned.md")

    assert "2004-04-05" in assigned
    assert "背景（不細寫）" in assigned
    assert "懸念／功能：核心問題" in assigned
    assert assigned.index("事件：") < assigned.index("懸念／功能：")
    assert "未分篇事件" not in assigned
    assert "未分篇事件" in unassigned
    assert "來源位置：" not in assigned
    assert "章節（ct）" not in assigned
    assert "累計（Ttl）" not in assigned


def test_transformation_counts_and_monthly_kpi_are_separate(tmp_path):
    documents = convert_outline_workbook.convert_workbook(
        _workbook(tmp_path / "outline.xlsx")
    )

    rendered = _generated(documents, "transformation_kpi.md")

    assert "| 2004-04 | 3 |" in rendered
    assert "| 2004-04 | 2004-04-05 | Operation Test | 主要事件 | 2 |" in rendered


def test_water_demon_colors_become_explicit_semantics(tmp_path):
    documents = convert_outline_workbook.convert_workbook(
        _workbook(tmp_path / "outline.xlsx")
    )

    rendered = _generated(documents, "water_demon_actions.md")

    assert "BBS（初登場）" in rendered
    assert "### 冰塊行動" in rendered
    assert rendered.index("## Siskin") < rendered.index("### 冰塊行動")
    assert "### 背景行動" in rendered
    assert "敘事層級：背景（不細寫）" in rendered
    assert "### 第7集｜星羅誕生" not in rendered
    assert "對應原著集數：第7集" in rendered
    assert "來源位置：" not in rendered


def test_actions_and_expenses_are_rendered_as_readable_sources(tmp_path):
    documents = convert_outline_workbook.convert_workbook(
        _workbook(tmp_path / "outline.xlsx")
    )

    action = _generated(documents, "operation_01.md")
    expenses = _generated(documents, "action_expenses.md")

    assert "日期：4月5-7日" in action
    assert "行動人員：Kafziel, Soar" in action
    assert "## 印度" in action
    assert "- 2, 11:00：離開檀香山" in action
    assert "| 機票 | 夏威夷-科契 | PREM | 1790 | 2 | 3580 |" in expenses


def test_filenames_are_english_and_indexed_in_source_order(tmp_path):
    documents = convert_outline_workbook.convert_workbook(
        _workbook(tmp_path / "outline.xlsx")
    )

    assert list(documents) == [
        "01_outline_operation_test.md",
        "02_outline_unassigned.md",
        "03_transformation_kpi.md",
        "04_water_demon_actions.md",
        "05_operation_01.md",
        "06_action_expenses.md",
    ]


def test_check_mode_detects_stale_output(tmp_path):
    documents = {"one.md": "# One\n"}
    output = tmp_path / "generated"

    assert not convert_outline_workbook.write_documents(documents, output)
    assert convert_outline_workbook.write_documents(documents, output, check=True)

    (output / "one.md").write_text("# Changed\n", encoding="utf-8")
    assert not convert_outline_workbook.write_documents(
        documents,
        output,
        check=True,
    )


def test_index_refresh_generates_sources_beside_workbook(tmp_path):
    _workbook(tmp_path / "大綱及行動.xlsx")

    _refresh_outline_sources(tmp_path)

    output = tmp_path / "00_大綱索引"
    assert (output / "01_outline_operation_test.md").exists()
    assert convert_outline_workbook.write_documents(
        convert_outline_workbook.convert_workbook(tmp_path / "大綱及行動.xlsx"),
        output,
        check=True,
    )
